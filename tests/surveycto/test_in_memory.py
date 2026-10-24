"""Dictionary, PII rules, stats and the workbook from in-memory objects."""

import io

import pytest
from openpyxl import load_workbook

from datasure.surveycto.cast import cast_wide
from datasure.surveycto.codebook import save_codebook
from datasure.surveycto.dictionary import build_dictionary
from datasure.surveycto.pii import PiiConfig, attach, classify, shows_values
from datasure.surveycto.questions import build_questions
from datasure.surveycto.stats import stats_rows
from datasure.surveycto.xml_contract import parse_contract
from tests.surveycto.farm_form import ROWS, SENTINELS, XML


@pytest.fixture
def built(tmp_path):
    path = tmp_path / "farm.xml"
    path.write_text(XML, encoding="utf-8")
    contract = parse_contract(path)
    questions = build_questions(contract)
    table = cast_wide(ROWS, contract)
    vardict = build_dictionary(table, questions, contract, SENTINELS)
    blocks = classify(table, vardict, questions, PiiConfig(), None)
    attach(vardict, blocks, questions, [])
    return table, questions, vardict, blocks


def test_rules_alone_leave_open_columns_flagged_unscanned(built):
    *_, blocks = built
    assert blocks["owner"] == {
        "status": "may_contain_pii",
        "basis": ["type:text", "unscanned"],
    }
    assert not shows_values(blocks["owner"]["status"])
    # The rules still decide what they can.
    assert blocks["village"]["status"] == "cleared_by_form"
    assert blocks["plots"]["status"] == "cleared_by_form"
    assert blocks["deviceid"]["status"] == "contains_pii"
    assert blocks["visit"]["status"] == "whitelisted"


def test_stats_rows_follow_skip_logic(built):
    table, questions, vardict, _ = built
    stats = {r["variable"]: r for r in stats_rows(table, questions, vardict, SENTINELS)}
    area = stats["area"]
    assert area["relevance"]
    assert area["eligible"] + area["skipped_by_logic"] == area["rows"]


def test_save_codebook_to_a_file_object(built):
    table, questions, vardict, _ = built
    buffer = io.BytesIO()
    save_codebook(buffer, vardict, stats_rows(table, questions, vardict, SENTINELS))

    workbook = load_workbook(io.BytesIO(buffer.getvalue()), read_only=True)
    assert workbook.sheetnames == ["Variables", "Summary stats"]


def test_parse_contract_reads_bytes(tmp_path):
    path = tmp_path / "farm.xml"
    path.write_text(XML, encoding="utf-8")
    assert (
        parse_contract(XML.encode()).nodes.keys() == parse_contract(path).nodes.keys()
    )
