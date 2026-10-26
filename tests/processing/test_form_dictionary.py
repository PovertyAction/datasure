"""Data dictionary and skip-logic stats built from an imported SurveyCTO form."""

import io
from unittest.mock import Mock, patch

import polars as pl
import pytest
from openpyxl import load_workbook

from datasure.connectors.scto import FormConfig, SurveyCTOClient
from datasure.processing.form_dictionary import (
    build_form_dictionary,
    dictionary_xlsx,
    form_table,
    sentinels_from_missing_codes,
)
from datasure.surveycto.cast import cast_wide
from datasure.surveycto.xml_contract import parse_contract
from datasure.utils.duckdb_utils import duckdb_get_table, load_form_definition
from tests.surveycto.farm_form import ROWS, XML

MISSING_CODES = pl.DataFrame(
    {"label": ["Don't know", "Refused", "Stata"], "codes": ["-99", "-88.0", ".d, -9.5"]}
)


@pytest.fixture(autouse=True)
def mock_database_functions():
    """Use real DuckDB files instead of the conftest mocks."""


@pytest.fixture
def project(tmp_path):
    """A project with the farm form imported from a mocked SurveyCTO server."""

    def cache_path(project_id, *parts):
        path = tmp_path.joinpath(project_id, *parts)
        path.parent.mkdir(parents=True, exist_ok=True)
        return path

    with patch("datasure.utils.duckdb_utils.get_cache_path", side_effect=cache_path):
        client = SurveyCTOClient("proj")
        client._scto_client = Mock()
        client._scto_client.list_datasets.return_value = []
        client._scto_client.download_form_xml.return_value = XML.encode()
        client._scto_client.download_form_data_json.return_value = ROWS
        client.import_data(
            FormConfig(alias="farm", form_id="farm", server="srv", username="a@b.org")
        )
        yield "proj"


@pytest.fixture
def built(project):
    raw = duckdb_get_table(project, "farm", "raw")
    return build_form_dictionary(project, "farm", raw, MISSING_CODES)


def test_sentinels_are_the_projects_whole_number_codes():
    assert sentinels_from_missing_codes(MISSING_CODES) == {
        -99: "Don't know",
        -88: "Refused",
    }


def test_a_code_under_two_labels_raises():
    codes = pl.DataFrame({"label": ["Don't know", "N/A"], "codes": ["-99", "-99, -97"]})
    with pytest.raises(ValueError, match="-99 is listed under both"):
        sentinels_from_missing_codes(codes)


def test_no_missing_codes_no_sentinels():
    empty = pl.DataFrame(schema={"label": pl.String, "codes": pl.String})
    assert sentinels_from_missing_codes(empty) == {}


def test_form_table_restores_the_cast_table(project):
    _, arrow_schema = load_form_definition(project, "farm")
    table = form_table(duckdb_get_table(project, "farm", "raw"), arrow_schema)

    cast = cast_wide(ROWS, parse_contract(XML.encode()))
    assert table.schema.equals(cast.schema, check_metadata=True)
    assert table.to_pylist() == cast.to_pylist()


def test_form_table_raises_when_raw_and_schema_differ(project):
    _, arrow_schema = load_form_definition(project, "farm")
    raw = duckdb_get_table(project, "farm", "raw").drop("area")
    with pytest.raises(ValueError, match=r"missing \['area'\]"):
        form_table(raw, arrow_schema)


def test_no_definition_no_dictionary(project):
    raw = duckdb_get_table(project, "farm", "raw")
    assert build_form_dictionary(project, "listing", raw, MISSING_CODES) is None


def test_stats_follow_skip_logic_and_set_codes_apart(built):
    stats = {row["variable"]: row for row in built.stats}
    # area is asked only when plots > 0; the row with plots = -99 skips it.
    assert stats["area"]["eligible"] == 1
    assert stats["area"]["skipped_by_logic"] == 1
    assert stats["plots"]["sentinel"] == 1
    assert "Don't know" in stats["plots"]["sentinel_detail"]


def test_pii_from_the_form_rules_trusting_publishable(built):
    status = {
        name: entry["pii"]["status"]
        for name, entry in built.dictionary["variables"].items()
    }
    assert status["village"] == "cleared_by_form"
    assert status["deviceid"] == "contains_pii"
    assert status["owner"] == "may_contain_pii"


def test_dictionary_workbook_hides_flagged_values(built):
    workbook = load_workbook(io.BytesIO(dictionary_xlsx(built)), read_only=True)
    assert workbook.sheetnames == ["Variables", "Summary stats"]
    cells = {
        str(cell)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    }
    assert not any("Ana Maria" in c or "3104567890" in c for c in cells)
